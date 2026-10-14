"""Genera tests/sample_tab_toc.docx: indice viejo estilo Word (titulo + TAB + pagina)
que mammoth convierte con espacios simples, y tests/sample_table_toc.docx (tabla 2 col)."""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def make_tab_toc(path):
    doc = Document()
    doc.add_heading("Indice", level=1)
    for title, page in [("Introduccion", 3), ("Marco teorico", 5), ("Conclusiones", 9)]:
        p = doc.add_paragraph()
        p.add_run(title)
        r = p.add_run("\t")
        p.add_run(str(page))
    doc.add_heading("Introduccion", level=1)
    doc.add_paragraph("Texto introductorio.")
    doc.add_heading("Marco teorico", level=1)
    doc.add_paragraph("Desarrollo.")
    doc.add_heading("Conclusiones", level=1)
    doc.add_paragraph("Cierre.")
    doc.save(path)

def make_table_toc(path):
    doc = Document()
    doc.add_heading("Indice", level=1)
    table = doc.add_table(rows=0, cols=2)
    for title, page in [("Introduccion", 3), ("Marco teorico", 5), ("Conclusiones", 9)]:
        row = table.add_row()
        row.cells[0].text = title
        row.cells[1].text = str(page)
    doc.add_heading("Introduccion", level=1)
    doc.add_paragraph("Texto introductorio.")
    doc.add_heading("Marco teorico", level=1)
    doc.add_paragraph("Desarrollo.")
    doc.add_heading("Conclusiones", level=1)
    doc.add_paragraph("Cierre.")
    doc.save(path)

make_tab_toc("tests/sample_tab_toc.docx")
make_table_toc("tests/sample_table_toc.docx")
print("ok")