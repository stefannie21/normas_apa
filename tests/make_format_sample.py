"""Genera tests/sample_format.docx: documento representativo de una tesis
con portada centrada, negritas, cursivas, justificado y sangria."""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

def para(text, align=None, bold=False, italic=False, size=12, indent=None, name="Normal"):
    p = doc.add_paragraph()
    if name:
        try:
            p.style = doc.styles[name]
        except Exception:
            pass
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.LEFT
    for i, seg in enumerate(text.split("**")):
        if seg == "":
            continue
        r = p.add_run(seg)
        r.bold = bold or (i % 2 == 1)
        r.italic = italic or (i % 2 == 1)
        r.font.size = Pt(size)
    if indent is not None:
        p.paragraph_format.first_line_indent = Pt(indent)
    return p

DOC = "text"
base_size = 12

para("UNIVERSIDAD NACIONAL DE PRUEBA", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14)
para("FACULTAD DE CIENCIAS", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
para("TITULO DE LA INVESTIGACION", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14)
para("Tesis presentada para optar al titulo de Licenciado", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)
para("Autor: Juan Perez", align=WD_ALIGN_PARAGRAPH.CENTER)
para("Tutor: Dra. Ana Gomez", align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()

doc.add_heading("Introduccion", level=1)
para(
    "Este es el primer parrafo del capitulo. Tiene varias lineas para que al paginar "
    "el texto pueda dividirse de forma **realista** entre dos hojas, conservando la "
    "negrita y la **cursiva** que el documento original declara, tal como la norma APA "
    "sugiere para citas y enfasis.",
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    indent=24.6,
)
para(
    "El segundo parrafo debe verse justificado a ambos margenes, con sangria de primera "
    "linea, y su **titulo entre comillas** se mantiene en negrita para representar un "
    "diario o fuente citada en texto.",
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    indent=24.6,
)

doc.add_heading("Marco teorico", level=1)
para(
    "Aqui se resume la teoria con citas. La referencia mas importante menciona a "
    "**Perez (2020)** y a **Gomez (2021)**, con elaboraciones propias en cursiva "
    "segun la muestra del texto proporcionado por el usuario.",
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    indent=24.6,
)

doc.add_heading("Conclusiones", level=1)
para(
    "Las conclusiones se presentan de manera **concisa** y organizada, retomando los "
    "objetivos de la investigacion y sugerencias para trabajos futuros.",
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    indent=24.6,
)

doc.save("tests/sample_format.docx")
print("ok")