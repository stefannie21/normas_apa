"""Genera tests/sample_static_toc.docx con un INDICE ESTATICO escrito a mano
(entradas con puntos y numeros de pagina), tipico de muchas tesis."""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

doc.add_heading("Indice", level=1)

entries = [
    ("Introduccion", 1),
    ("Marco teorico", 2),
    ("Metodologia", 3),
    ("Conclusiones", 4),
]
for title, page in entries:
    p = doc.add_paragraph()
    r = p.add_run(title + "...................................................")
    r2 = p.add_run(str(page))
    r2.bold = False
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

doc.add_page_break()
doc.add_heading("Introduccion", level=1)
doc.add_paragraph("Texto de la introduccion de la tesis.")
doc.add_heading("Marco teorico", level=1)
doc.add_paragraph("Desarrollo teorico.")
doc.add_heading("Metodologia", level=2)
doc.add_paragraph("Descripcion de la metodologia.")
doc.add_heading("Conclusiones", level=1)
doc.add_paragraph("Conclusiones del trabajo.")

doc.save("tests/sample_static_toc.docx")
print("ok")