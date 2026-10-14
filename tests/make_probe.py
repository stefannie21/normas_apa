"""Sonda: crea un docx con salto de pagina, interlineado y seccion apaisada,
y muestra que conserva el XML original."""
import re
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.shared import Pt

doc = Document()
doc.add_heading("Uno", level=1)

p = doc.add_paragraph("Parrafo con interlineado doble.")
p.paragraph_format.line_spacing = 2.0
p.paragraph_format.space_after = Pt(12)

doc.add_paragraph("Antes del salto.")
doc.add_page_break()
doc.add_paragraph("Despues del salto de pagina.")

sec = doc.add_section(WD_SECTION.NEW_PAGE)
sec.orientation = WD_ORIENT.LANDSCAPE
doc.add_paragraph("Seccion apaisada.")

doc.save("tests/sample_probe.docx")

xml = zipfile.ZipFile("tests/sample_probe.docx").read("word/document.xml").decode("utf8")
print("pgSz:", re.findall(r"<w:pgSz[^>]*/>", xml))
print("orient:", re.findall(r'w:orient="[a-z]+"', xml))
print("saltos de pagina:", xml.count('w:type="page"'))
print("spacing:", re.findall(r"<w:spacing[^>]*/>", xml)[:4])