"""Comprueba que el score de cumplimiento sea sensato:
documento limpio alto y documento con errores puntuales casi igual de alto."""
import sys
from io import BytesIO

sys.path.insert(0, ".")

from docx import Document
from docx.shared import Inches, Pt

from backend.app.services.analyzer import analyze
from backend.app.services.docx_parser import parse_docx


def build(n_paragraphs: int, break_one: bool = False) -> bytes:
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    for i in range(n_paragraphs):
        p = doc.add_paragraph(
            "Parrafo de prueba numero %d con texto suficiente para contar como cuerpo "
            "del documento y comprobar el nivel de cumplimiento (Author, 2020). "
            "Este documento no tiene otro problema que el intencional." % (i + 1)
        )
        p.paragraph_format.first_line_indent = Inches(0.5)
        if break_one and i == n_paragraphs // 2:
            p.paragraph_format.line_spacing = 1.0
    doc.add_heading("Referencias", level=1)
    ref = doc.add_paragraph()
    ref.paragraph_format.left_indent = Inches(0.5)
    ref.paragraph_format.first_line_indent = Inches(-0.5)
    ref.add_run("Author, A. (2020). ")
    r = ref.add_run("Titulo del trabajo")
    r.italic = True
    ref.add_run(". Editorial.")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def score(data: bytes) -> dict:
    return analyze(parse_docx(data, "score.docx")).summary


def main() -> int:
    a = score(build(200))
    b = score(build(500, break_one=True))
    print("200 parrafos, documento limpio :", a["compliance_score"], "% |",
          a["total_findings"], "hallazgos")
    print("500 parrafos, 1 solo error     :", b["compliance_score"], "% |",
          b["total_findings"], "hallazgos")
    assert a["compliance_score"] >= 95, f"limpio deberia puntuar alto: {a['compliance_score']}"
    assert b["compliance_score"] >= 90, f"un error puntual no debe hundir la nota: {b['compliance_score']}"
    print("OK: el score responde a la realidad del documento")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())