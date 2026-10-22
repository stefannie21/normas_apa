"""Pruebas de validación estricta para la Configuración Global APA 7:
1. Ámbito global de aplicación.
2. Apellidos compuestos y nombres en español (dobles apellidos, fuzzy matching antes de iniciales, no falso CRUCE-001).
3. Autores institucionales / corporativos (SUNAT, OECD, Municipalidad Metropolitana de Lima, El Comercio, etc., prevención de REF-007).
4. Detección flexible de año de publicación (AAAA) en autoría/fecha (prevención de REF-004).
"""
import sys
from pathlib import Path
from io import BytesIO

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.shared import Inches, Pt

from backend.app.services.analyzer import analyze
from backend.app.services.docx_parser import parse_docx
from backend.app.services.bibliography import (
    author_matches, is_corporate_author, author_prefix_before_initials, extract_citations
)


def test_apellidos_compuestos():
    """Criterio 2: Dobles apellidos en español como un único bloque y concordancia flexible."""
    print("-> Probando regla de apellidos compuestos y fuzzy matching...")

    # Documento con citas parentéticas y narrativas con dobles apellidos
    doc = Document()
    p1 = doc.add_paragraph("Según señalan Estrada Frías y Chávez Gamarra (2023), los resultados son consistentes.")
    p2 = doc.add_paragraph("Asimismo, en investigaciones previas (Limaco García, 2020) se encontraron efectos similares.")
    p3 = doc.add_paragraph("Como argumenta Limaco García et al. (2020), la teoría sigue vigente.")

    doc.add_heading("Referencias", level=1)
    # Referencias con diferentes variantes de formato antes de las iniciales
    ref1 = doc.add_paragraph("Estrada Frías, J. A., y Chávez Gamarra, C. (2023). Modelos de evaluación educativa. Editorial Universitaria.")
    ref2 = doc.add_paragraph("Limaco García, M. (2020). Innovación y aprendizaje en educación superior. Fondo Editorial.")

    for r in [ref1, ref2]:
        r.paragraph_format.left_indent = Inches(0.5)
        r.paragraph_format.first_line_indent = Inches(-0.5)

    buf = BytesIO()
    doc.save(buf)

    parsed = parse_docx(buf.getvalue(), "test_compuestos.docx")
    result = analyze(parsed)
    findings = result.findings
    bib = parsed.bibliography

    # Verificar que las citas capturaron los dobles apellidos como bloques únicos
    all_surnames = [s for c in bib.citations for s in c.surnames]
    assert "Estrada Frías" in all_surnames, f"Debe identificar 'Estrada Frías': {all_surnames}"
    assert "Chávez Gamarra" in all_surnames, f"Debe identificar 'Chávez Gamarra': {all_surnames}"
    assert "Limaco García" in all_surnames, f"Debe identificar 'Limaco García': {all_surnames}"

    # Prohibido marcar CRUCE-001 (cita sin referencia) para estas citas
    cruce_findings = [f for f in findings if f.rule_id == "CRUCE-001"]
    assert len(cruce_findings) == 0, f"No debe haber citas sin referencia para apellidos compuestos válidos: {[f.message for f in cruce_findings]}"
    print("   [OK] Dobles apellidos y concordancia flexible validados sin falsos positivos.")


def test_autores_institucionales_ref007():
    """Criterio 3: Reconocimiento de autores institucionales y prevención de REF-007."""
    print("-> Probando autores institucionales / corporativos (prevención de REF-007)...")

    doc = Document()
    doc.add_paragraph("De acuerdo con la SUNAT (2021), la recaudación aumentó.")
    doc.add_paragraph("La OECD (2019) reporta crecimiento económico moderado.")
    doc.add_paragraph("La Municipalidad Metropolitana de Lima (2022) publicó su plan de acción.")
    doc.add_paragraph("Según informó El Comercio (2020), la cobertura se expandió.")

    doc.add_heading("Referencias", level=1)
    r1 = doc.add_paragraph("SUNAT. (2021). Informe de recaudación tributaria nacional. Superintendencia Nacional de Aduanas y de Administración Tributaria.")
    r2 = doc.add_paragraph("OECD. (2019). OECD Economic Outlook. OECD Publishing.")
    r3 = doc.add_paragraph("Municipalidad Metropolitana de Lima. (2022). Plan de desarrollo local concertado 2022-2030. MML.")
    r4 = doc.add_paragraph("El Comercio. (2020). Análisis del impacto socioeconómico nacional. Diario El Comercio.")

    for r in [r1, r2, r3, r4]:
        r.paragraph_format.left_indent = Inches(0.5)
        r.paragraph_format.first_line_indent = Inches(-0.5)

    buf = BytesIO()
    doc.save(buf)

    parsed = parse_docx(buf.getvalue(), "test_institucionales.docx")
    result = analyze(parsed)

    # Verificar que NO se emite REF-007 ("Referencia sin autor reconocible")
    ref007_findings = [f for f in result.findings if f.rule_id == "REF-007" and "sin autor" in f.message.lower()]
    assert len(ref007_findings) == 0, f"Autores corporativos no deben generar REF-007: {[f.message for f in ref007_findings]}"

    # Verificar que NO se emite CRUCE-001 para autores institucionales coincidentes
    cruce_findings = [f for f in result.findings if f.rule_id == "CRUCE-001"]
    assert len(cruce_findings) == 0, f"No debe haber CRUCE-001 para citas corporativas presentes: {[f.message for f in cruce_findings]}"
    print("   [OK] Autores institucionales reconocidos y exentos de REF-007.")


def test_deteccion_flexible_ano_ref004():
    """Criterio 4: Detección flexible de año (AAAA) en autoría/fecha y prevención de REF-004."""
    print("-> Probando detección flexible de año en referencias (prevención de REF-004)...")

    doc = Document()
    doc.add_paragraph("La metodología se basa en diversos enfoques clásicos y contemporáneos.")

    doc.add_heading("Referencias", level=1)
    # Variantes de año en el encabezado
    r1 = doc.add_paragraph("García, M. (2020). Metodología cuantitativa avanzada. Editorial Tecnos.")
    r2 = doc.add_paragraph("López, R. ( 2021 ). Evaluación formativa del aprendizaje. Editorial Morata.")
    r3 = doc.add_paragraph("Martínez, C. (2022a). Teorías del aprendizaje situado. Alianza Editorial.")
    r4 = doc.add_paragraph("Sánchez, P. (2023, 15 de marzo). Transformación digital. McGraw-Hill.")
    r5 = doc.add_paragraph("Ministerio de Salud. (s.f.). Guía técnica de intervención comunitaria. MINSA.")

    for r in [r1, r2, r3, r4, r5]:
        r.paragraph_format.left_indent = Inches(0.5)
        r.paragraph_format.first_line_indent = Inches(-0.5)

    buf = BytesIO()
    doc.save(buf)

    parsed = parse_docx(buf.getvalue(), "test_anos.docx")
    result = analyze(parsed)

    # Verificar que NINGUNA de estas referencias dispare REF-004 ("Referencia sin año de publicación")
    ref004_findings = [f for f in result.findings if f.rule_id == "REF-004"]
    assert len(ref004_findings) == 0, f"No deben emitir REF-004: {[f.message for f in ref004_findings]}"
    print("   [OK] Detección flexible de año previene correctamente REF-004.")


if __name__ == "__main__":
    test_apellidos_compuestos()
    test_autores_institucionales_ref007()
    test_deteccion_flexible_ano_ref004()
    print("\nTODAS LAS PRUEBAS DE CONFIGURACIÓN GLOBAL APA 7 PASARON EXITOSAMENTE.")
