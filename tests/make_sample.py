"""Genera un .docx de ejemplo con incumplimientos APA intencionales para pruebas."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def build_sample() -> bytes:
    doc = Document()

    # Margenes incorrectos (0.7 in) y fuente incorrecta (Courier New 13).
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Courier New"
    normal.font.size = Pt(13)

    # Titulo (nivel 1) que no esta centrado.
    h1 = doc.add_heading("Introduccion", level=1)
    h1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in h1.runs:
        run.font.name = "Courier New"
        run.font.size = Pt(13)

    # Parrafo justificado, interlineado sencillo, sin sangria.
    p = doc.add_paragraph(
        "La investigacion academica requiere rigor metodologico y una correcta "
        "citacion de las fuentes utilizadas en el desarrollo del proyecto."
    )
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Inches(0.0)
    p.paragraph_format.space_after = Pt(10)

    # Cita sin ano y con 'and' dentro del parentesis.
    p2 = doc.add_paragraph(
        "Diversos autores senalan la importancia del metodo (Garcia and Lopez) "
        "y tambien hay estudios previos sin ano."
    )
    p2.paragraph_format.line_spacing = 1.0

    # Salto de jerarquia: nivel 1 -> nivel 3.
    doc.add_heading("Antecedentes", level=3)

    # Cita textual sin pagina.
    p3 = doc.add_paragraph('Segun el autor, "el exito depende del esfuerzo" (Garcia, 2020).')
    p3.paragraph_format.line_spacing = 1.0

    # Seccion de referencias con orden incorrecto y DOI en formato viejo.
    doc.add_heading("Referencias", level=1)
    ref1 = doc.add_paragraph("Zapata, L. (2019). Estudios sobre metodo. Editorial Norte.")
    ref1.paragraph_format.line_spacing = 1.0
    ref2 = doc.add_paragraph("Alvarez, M. (2018). Fundamentos de investigacion. doi:10.1000/xyz123")
    ref2.paragraph_format.line_spacing = 1.0

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def write_sample(path: Path) -> Path:
    path.write_bytes(build_sample())
    return path


if __name__ == "__main__":
    out = Path(__file__).with_name("sample_tesis.docx")
    write_sample(out)
    print(f"Documento generado: {out}")
