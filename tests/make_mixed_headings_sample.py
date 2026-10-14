"""Genera tests/sample_mixed_headings.docx con titulos CON estilo nativo y SIN estilo
para comprobar HEAD-004 (deteccion + ubicacion)."""
from docx import Document
from docx.shared import Pt

doc = Document()

# --- Con estilo nativo (correcto, no debe reportarse) ---
doc.add_heading("Con estilo nativo", level=1)
doc.add_paragraph("Parrafo de cuerpo del documento para tener contenido.")

# --- Sin estilo: mayusculas + negrita (debe reportarse) ---
p = doc.add_paragraph()
r = p.add_run("PROBLEMA DE INVESTIGACIÓN")
r.bold = True
r.font.size = Pt(14)
doc.add_paragraph("Descripcion del problema de investigacion con varias lineas de texto normal para el cuerpo.")

# --- Sin estilo: solo negrita ---
p = doc.add_paragraph()
r = p.add_run("Objetivo general")
r.bold = True
doc.add_paragraph("El objetivo general del trabajo de investigacion.")

# --- Sin estilo: numeracion de capitulo ---
p = doc.add_paragraph()
r = p.add_run("2.1. Justificación de la investigación")
r.bold = True
r.font.size = Pt(13)
doc.add_paragraph("Texto de justificación.")

# --- Texto normal (NO debe reportarse) ---
doc.add_paragraph("Este es un parrafo de cuerpo normal, sin negrita ni mayusculas, que no debe ser marcado como titulo.")

doc.save("tests/sample_mixed_headings.docx")
print("ok")