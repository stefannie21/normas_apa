"""Revision de orientacion de secciones/paginas en un .docx.

Uso:
    python tools/check_orientation.py "ruta/al/documento.docx"
"""
from __future__ import annotations

import sys
from collections import Counter

from docx import Document
from docx.enum.section import WD_ORIENT

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
TWIPS_PER_IN = 1440.0


def _twips_to_in(val, default):
    if val is None:
        return default
    try:
        return round(float(val) / TWIPS_PER_IN, 3)
    except (TypeError, ValueError):
        return default


def collect_sectpr(doc):
    sect_prs = []
    try:
        body = doc.element.body
        for child in body.iterchildren():
            if child.tag == W_NS + "p":
                ppr = child.find(W_NS + "pPr")
                if ppr is not None:
                    sp = ppr.find(W_NS + "sectPr")
                    if sp is not None:
                        sect_prs.append(sp)
            elif child.tag == W_NS + "sectPr":
                sect_prs.append(child)
    except Exception:
        sect_prs = []
    if not sect_prs:
        try:
            sect_prs = [s._sectPr for s in doc.sections]
        except Exception:
            sect_prs = []
    return sect_prs


def section_info(sp, index):
    pg_sz = sp.find(W_NS + "pgSz")
    orient_attr = None
    w_twips = None
    h_twips = None
    if pg_sz is not None:
        orient_attr = pg_sz.get(W_NS + "orient")
        w_twips = pg_sz.get(W_NS + "w")
        h_twips = pg_sz.get(W_NS + "h")

    w = _twips_to_in(w_twips, 8.268)
    h = _twips_to_in(h_twips, 11.693)
    if w <= 0:
        w = 8.268
    if h <= 0:
        h = 11.693

    landscape = orient_attr == "landscape" or w > h
    if landscape and w < h:
        w, h = h, w

    pg_mar = sp.find(W_NS + "pgMar")
    top = _twips_to_in(pg_mar.get(W_NS + "top") if pg_mar is not None else None, 1.0)
    bottom = _twips_to_in(pg_mar.get(W_NS + "bottom") if pg_mar is not None else None, 1.0)
    left = _twips_to_in(pg_mar.get(W_NS + "left") if pg_mar is not None else None, 1.0)
    right = _twips_to_in(pg_mar.get(W_NS + "right") if pg_mar is not None else None, 1.0)

    return {
        "index": index,
        "orientation": "HORIZONTAL (landscape)" if landscape else "VERTICAL (portrait)",
        "landscape": landscape,
        "width_in": w,
        "height_in": h,
        "margins_in": (top, bottom, left, right),
        "orient_attr": orient_attr,
        "wd_orient": WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT,
    }


def map_paragraphs_to_sections(doc):
    section_of = []
    idx = 0
    for p in doc.paragraphs:
        section_of.append(idx)
        try:
            ppr = p._p.pPr
            if ppr is not None and ppr.find(W_NS + "sectPr") is not None:
                idx += 1
        except Exception:
            pass
    return section_of


def first_key_per_section(doc, section_of):
    first_key = {}
    for i, p in enumerate(doc.paragraphs):
        sec = section_of[i] if i < len(section_of) else 0
        key = (p.text or "").strip().lower()[:48]
        if key and sec not in first_key:
            first_key[sec] = key
    return first_key


def estimate_pages_per_section(doc, section_of, num_sections):
    pages = [1] * num_sections
    for i, p in enumerate(doc.paragraphs):
        sec = section_of[i] if i < len(section_of) else 0
        xml = p._p.xml
        if "w:br" in xml and 'type="page"' in xml:
            pages[sec] += 1
        if "lastRenderedPageBreak" in xml:
            pages[sec] += 1
    return pages


def main(path):
    doc = Document(path)
    sect_prs = collect_sectpr(doc)
    sections = [section_info(sp, i) for i, sp in enumerate(sect_prs)]
    section_of = map_paragraphs_to_sections(doc)
    first_keys = first_key_per_section(doc, section_of)
    page_counts = estimate_pages_per_section(doc, section_of, len(sections))

    print("=" * 78)
    print("REPORTE DE ORIENTACION DE SECCIONES/PAGINAS")
    print("Documento:", path)
    print("Total de secciones detectadas:", len(sections))
    print("Total de parrafos:", len(doc.paragraphs))
    print("=" * 78)
    print(f"{'Sec':>4}  {'Orientacion':<26} {'Ancho(in)':>10} {'Alto(in)':>9}  Margenes(T/B/L/R)   Inicio")
    print("-" * 78)
    for s in sections:
        i = s["index"]
        fk = first_keys.get(i, "(sin texto)")
        margins = s["margins_in"]
        print(
            f"{i + 1:>4}  {s['orientation']:<26} {s['width_in']:>10} {s['height_in']:>9}  "
            f"{margins[0]}/{margins[1]}/{margins[2]}/{margins[3]}  {fk!r}"
        )

    print()
    c = Counter(s["landscape"] for s in sections)
    print("Resumen: verticales =", c.get(False, 0), "| horizontales =", c.get(True, 0))
    print()
    print("Secciones HORIZONTALES:")
    for s in sections:
        if s["landscape"]:
            print(f"  - Seccion {s['index'] + 1}: {s['width_in']} x {s['height_in']} in  (primer parrafo: {first_keys.get(s['index'], '')!r})")
    print()
    print("Secciones VERTICALES:")
    for s in sections:
        if not s["landscape"]:
            print(f"  - Seccion {s['index'] + 1}: {s['width_in']} x {s['height_in']} in")
    print()
    print("Estimacion de paginas por seccion (saltos de pagina manuales):")
    for i, n in enumerate(page_counts):
        orient = "H" if sections[i]["landscape"] else "V"
        print(f"  Seccion {i + 1} [{orient}]: ~{n} pagina(s)")
    print("=" * 78)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uso: python tools/check_orientation.py <archivo.docx>")
        sys.exit(1)
    main(sys.argv[1])
